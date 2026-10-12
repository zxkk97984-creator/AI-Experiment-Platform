"""Decimal weighted grades with explicit missing/exemption decisions and frozen publications."""
from copy import deepcopy
import csv
from decimal import Decimal, ROUND_HALF_UP
import hashlib
from io import BytesIO, StringIO
import json
from uuid import uuid4

from openpyxl import Workbook
from fastapi import HTTPException
from sqlalchemy import delete, select, update

from app.errors import api_error
from app.models import (AcademicTerm, Assignment, CodeGrade, Course, Exam, ExamQuestion, ExamSubmission,
                        ExperimentModule, ExperimentRecord, ExperimentSubmission, JudgeQuestion,
                        LabGradeRevision, LabRecord, LabRun, LabSubmission, Lesson, Submission, TeachingClass, User)
from app.models.gradebook import (Gradebook, GradebookAdjustment, GradebookCategory,
                                 GradebookColumn, GradebookEntry, GradebookSnapshot)
from app.models.teaching import CourseOffering, OfferingEnrollment, OfferingExperiment
from app.services.account_service import audit, lock_policy
from app.services.offering_service import require_offering
from app.services.time_utils import utc_now
from app.services.gradebook_sources import preload_sources

HUNDRED = Decimal(100)


def decimal(value):
    result = Decimal(str(value))
    if not result.is_finite(): raise api_error(422, 'INVALID_SOURCE_SCORE', '成绩来源含非有限数值')
    return result


def score_text(value):
    return str(value.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)) if value is not None else None


def default_config():
    return [{'key': key, 'title': title, 'weight': str(weight), 'columns': [
        {'key': f'{key}-main', 'title': title, 'weight': '100', 'source_kind': 'manual', 'source_id': None, 'selection': 'highest'}]}
        for key, title, weight in [('regular', '平时成绩', 30), ('labs', '实验成绩', 40), ('exam', '考核成绩', 30)]]


def book_for(db, offering_id, *, lock=False):
    query = select(Gradebook).where(Gradebook.offering_id == offering_id)
    return db.scalar(query.with_for_update().execution_options(populate_existing=True) if lock else query)


def lock_authority(db, offering_id, actor, *, write=True):
    # Shared with roster/account changes and short lab score commits. Acquire
    # this before any actor/offering/source row: never hold it during execution.
    lock_policy(db)
    authenticated_version = actor.session_version
    current = db.scalar(select(User).where(User.id == actor.id).with_for_update()
                        .execution_options(populate_existing=True))
    if (not current or current.status != 'active' or current.role not in ('teacher', 'admin')
            or current.session_version != authenticated_version):
        raise api_error(403, 'FORBIDDEN', '当前账号已失去成绩管理权限，请重新登录')
    offering = db.scalar(select(CourseOffering).where(CourseOffering.id == offering_id).with_for_update()
                         .execution_options(populate_existing=True))
    if not offering: raise api_error(404, 'OFFERING_NOT_FOUND', '授课实例不存在')
    term = db.scalar(select(AcademicTerm).where(AcademicTerm.id == offering.academic_term_id)
                     .with_for_update().execution_options(populate_existing=True))
    klass = db.scalar(select(TeachingClass).where(TeachingClass.id == offering.teaching_class_id)
                      .with_for_update().execution_options(populate_existing=True))
    if write and (not term or not klass):
        raise api_error(409, 'OFFERING_READ_ONLY', '授课实例、班级或学期已归档，当前只读')
    # require_offering now reads the objects refreshed by current locking reads.
    # expire + a normal SELECT would still read an old REPEATABLE READ snapshot.
    return require_offering(db, offering_id, current, write=write)


def lock_book(db, offering_id, actor, expected_revision, *, write=True, create=False):
    lock_authority(db, offering_id, actor, write=write)
    book = book_for(db, offering_id, lock=True)
    if book is None:
        if not create or expected_revision != 0:
            raise api_error(409, 'GRADEBOOK_REVISION_CONFLICT', '成绩簿尚未配置或版本已变化，请刷新')
        book = Gradebook(offering_id=offering_id, revision=1)
        db.add(book); db.flush()
    else:
        result = db.execute(update(Gradebook).where(Gradebook.id == book.id,
                            Gradebook.revision == expected_revision).values(revision=Gradebook.revision + 1))
        if result.rowcount != 1:
            raise api_error(409, 'GRADEBOOK_REVISION_CONFLICT', '成绩簿已被其他操作修改，请刷新后重试')
        book = book_for(db, offering_id, lock=True)
    return book


def categories(db, book, *, lock=False):
    rows = query_rows(db, select(GradebookCategory).where(GradebookCategory.gradebook_id == book.id)
                      .order_by(GradebookCategory.position, GradebookCategory.id), lock)
    columns = query_rows(db, select(GradebookColumn).where(GradebookColumn.gradebook_id == book.id)
                         .order_by(GradebookColumn.position, GradebookColumn.id), lock)
    return rows, columns


def config_read(db, book, *, lock=False):
    groups, columns = categories(db, book, lock=lock)
    return [{'key': c.key, 'title': c.title, 'weight': str(c.weight), 'columns': [
        {'key': col.key, 'title': col.title, 'weight': str(col.weight), 'source_kind': col.source_kind,
         'source_id': col.source_id, 'selection': col.selection} for col in columns if col.category_id == c.id]} for c in groups]


def source_target(db, kind, ident, offering_id, *, lock=False):
    model = {'assignment': Assignment, 'exam': Exam, 'experiment': OfferingExperiment}.get(kind)
    if not model: return None
    query = select(model).where(model.id == ident)
    if lock: query = query.with_for_update().execution_options(populate_existing=True)
    source = db.scalar(query)
    if not source or source.offering_id != offering_id:
        raise api_error(422, 'GRADE_SOURCE_SCOPE', '成绩来源必须属于当前授课实例')
    if source.status != 'published':
        raise api_error(422, 'GRADE_SOURCE_UNPUBLISHED', '只能绑定已发布的班级教学活动')
    return source


def save_config(db, offering_id, actor, payload):
    book = lock_book(db, offering_id, actor, payload.expected_revision, create=True)
    for category in payload.categories:
        for column in category.columns:
            if column.source_kind != 'manual': source_target(db, column.source_kind, column.source_id, offering_id, lock=True)
    old_categories, old_columns = categories(db, book, lock=True)
    category_map = {row.key: row for row in old_categories}
    column_map = {row.key: row for row in old_columns}
    keep_categories, keep_columns = set(), set()
    for pos, cfg in enumerate(payload.categories):
        category = category_map.get(cfg.key)
        if not category:
            category = GradebookCategory(gradebook_id=book.id, key=cfg.key)
            db.add(category)
        category.title, category.weight, category.position = cfg.title, cfg.weight, pos
        db.flush(); keep_categories.add(category.id)
        for index, col_cfg in enumerate(cfg.columns):
            column = column_map.get(col_cfg.key)
            if not column:
                column = GradebookColumn(gradebook_id=book.id, key=col_cfg.key)
                db.add(column)
            elif (column.source_kind, column.source_id) != (col_cfg.source_kind, col_cfg.source_id):
                db.execute(delete(GradebookEntry).where(GradebookEntry.column_id == column.id))
            column.category_id, column.title, column.weight = category.id, col_cfg.title, col_cfg.weight
            column.source_kind, column.source_id = col_cfg.source_kind, col_cfg.source_id
            column.selection, column.position = col_cfg.selection, index
            db.flush(); keep_columns.add(column.id)
    removed_columns = [c.id for c in old_columns if c.id not in keep_columns]
    if removed_columns:
        db.execute(delete(GradebookEntry).where(GradebookEntry.column_id.in_(removed_columns)))
        db.execute(delete(GradebookColumn).where(GradebookColumn.id.in_(removed_columns)))
    db.execute(delete(GradebookCategory).where(GradebookCategory.gradebook_id == book.id,
                                               GradebookCategory.id.not_in(keep_categories)))
    audit(db, actor.id, 'gradebook.configure', 'offering', offering_id, changes={'revision': book.revision})
    return book


def roster(db, offering_id, *, lock=False):
    # Lock membership and users explicitly. A normal joined SELECT could reuse
    # both an old InnoDB snapshot and already cached ORM identities.
    members = query_rows(db, select(OfferingEnrollment).where(OfferingEnrollment.offering_id == offering_id,
        OfferingEnrollment.status == 'active').order_by(OfferingEnrollment.student_id), lock)
    return query_rows(db, select(User).where(User.id.in_([m.student_id for m in members]), User.role == 'student')
                      .order_by(User.student_no, User.id), lock)


def require_student(db, offering_id, student_id, *, lock=False):
    query = select(OfferingEnrollment.id).join(User, User.id == OfferingEnrollment.student_id).where(
        OfferingEnrollment.offering_id == offering_id, OfferingEnrollment.student_id == student_id,
        OfferingEnrollment.status == 'active', User.role == 'student')
    found = db.scalar(query.with_for_update().execution_options(populate_existing=True) if lock else query)
    if found is None: raise api_error(422, 'GRADEBOOK_STUDENT_SCOPE', '学生不在当前授课实例')


def update_entries(db, offering_id, actor, payload):
    book = lock_book(db, offering_id, actor, payload.expected_revision)
    columns = {c.key: c for c in categories(db, book, lock=True)[1]}
    for change in payload.entries:
        require_student(db, offering_id, change.student_id, lock=True)
        column = columns.get(change.column_key)
        if not column: raise api_error(422, 'GRADEBOOK_COLUMN_NOT_FOUND', '成绩栏目不存在')
        if change.state == 'auto' and column.source_kind == 'manual':
            raise api_error(422, 'GRADEBOOK_SOURCE_REQUIRED', '人工栏目不能自动取分')
        if change.state == 'scored' and column.source_kind != 'manual' and not (change.reason and change.reason.strip()):
            raise api_error(422, 'GRADEBOOK_REASON_REQUIRED', '覆盖来源成绩必须填写调分原因')
        entry = db.scalar(select(GradebookEntry).where(GradebookEntry.column_id == column.id,
                                                       GradebookEntry.student_id == change.student_id)
                          .with_for_update().execution_options(populate_existing=True))
        if entry is None:
            entry = GradebookEntry(column_id=column.id, student_id=change.student_id)
            db.add(entry)
        before = {'state': entry.state, 'score': score_text(entry.score)}
        entry.state, entry.score, entry.reason = change.state, change.score, change.reason.strip() if change.reason else None
        entry.updated_by_id = actor.id
        audit(db, actor.id, 'gradebook.entry', 'gradebook', book.id, reason=entry.reason,
              changes={'student_id': change.student_id, 'column': column.key, 'before': before,
                       'state': change.state, 'score': score_text(change.score), 'revision': book.revision})
    db.flush()
    return book


def adjust(db, offering_id, actor, payload):
    book = lock_book(db, offering_id, actor, payload.expected_revision)
    require_student(db, offering_id, payload.student_id, lock=True)
    current = db.scalar(select(GradebookAdjustment).where(GradebookAdjustment.gradebook_id == book.id,
                                                         GradebookAdjustment.student_id == payload.student_id)
                            .with_for_update().execution_options(populate_existing=True))
    previous = score_text(current.score) if current else None
    if payload.score is None:
        if current: db.delete(current)
    else:
        if current is None:
            current = GradebookAdjustment(gradebook_id=book.id, student_id=payload.student_id)
            db.add(current)
        current.score, current.reason, current.updated_by_id = payload.score, payload.reason, actor.id
    audit(db, actor.id, 'gradebook.adjust', 'gradebook', book.id, reason=payload.reason,
          changes={'student_id': payload.student_id, 'before': previous, 'after': score_text(payload.score), 'revision': book.revision})
    db.flush()
    return book


def pending(message, state='pending'):
    return {'state': state, 'score': None, 'message': message, 'source': None}


def source_result(score, evidence):
    value = decimal(score)
    if value < 0 or value > HUNDRED: raise api_error(422, 'INVALID_SOURCE_SCORE', '来源百分制成绩超出0至100范围')
    return {'state': 'scored', 'score': score_text(value), 'message': '', 'source': evidence}


def query_rows(db, query, lock):
    return db.scalars(query.with_for_update().execution_options(populate_existing=True) if lock else query).all()


def select_score(rows, selection):
    return max(rows, key=lambda row: (decimal(row.score), row.id)) if selection == 'highest' else max(rows, key=lambda row: row.id)


def source_score(db, column, offering_id, student_id, *, lock=False, missing_confirmed=False, target=None, source_rows=None):
    def fetch(bucket, key, query):
        return source_rows[bucket].get(key, []) if source_rows is not None else query_rows(db, query, lock)

    try:
        target = target or source_target(db, column.source_kind, column.source_id, offering_id, lock=lock)
    except HTTPException as exc:
        if exc.detail['code'] in ('GRADE_SOURCE_SCOPE', 'GRADE_SOURCE_UNPUBLISHED'):
            return pending('原成绩来源不可用，请重新确认分项；官方历史版本保留')
        raise
    if column.source_kind == 'assignment':
        questions = fetch('questions', target.id, select(JudgeQuestion).where(JudgeQuestion.assignment_id == target.id).order_by(JudgeQuestion.id))
        if not questions: return pending('作业尚无有效题目')
        earned, full, evidence = Decimal(0), Decimal(0), []
        for question in questions:
            maximum = decimal(getattr(question, 'max_score', 100))
            if maximum <= 0: return pending('题目满分必须大于0')
            rows = fetch('submissions', (question.id, student_id), select(Submission).where(Submission.question_id == question.id, Submission.student_id == student_id).order_by(Submission.id))
            if not rows:
                if not missing_confirmed: return pending('存在未提交题目，需要教师确认缺交', 'missing_pending')
                full += maximum
                evidence.append({'question_id': question.id, 'score': '0', 'max_score': str(maximum), 'confirmed_missing': True})
                continue
            if any(r.grading_status != 'completed' or r.score is None for r in rows): return pending('作业提交仍在评分或需要处理')
            valid, grades = [], {}
            for row in rows:
                grade_rows = fetch('code_grades', row.id, select(CodeGrade).where(CodeGrade.submission_id == row.id).order_by(CodeGrade.id))
                grade = grade_rows[0] if grade_rows else None
                if grade and grade.mode == 'active' and (grade.status != 'completed' or grade.needs_teacher_review):
                    return pending('作业AI评分尚未公开完成')
                grades[row.id] = grade
                valid.append(row)
            chosen = select_score(valid, column.selection)
            if not Decimal(0) <= decimal(chosen.score) <= maximum:
                raise api_error(422, 'INVALID_SOURCE_SCORE', '题目来源分数超过该题满分')
            earned += decimal(chosen.score); full += maximum
            chosen_grade = grades[chosen.id]
            evidence.append({'question_id': question.id, 'submission_id': chosen.id, 'score': str(chosen.score),
                             'max_score': str(maximum), 'updated_at': str(chosen.updated_at),
                             'code_grade_id': chosen_grade.id if chosen_grade else None,
                             'rubric_id': chosen_grade.rubric_id if chosen_grade else None})
        return source_result(earned / full * HUNDRED, {'kind': 'assignment', 'id': target.id, 'selection': column.selection,
                             'task_updated_at': str(target.updated_at), 'questions': evidence})
    if column.source_kind == 'exam':
        if not (target.show_score_after_grading or target.review_released_at): return pending('考试成绩尚未公开')
        rows = fetch('exams', (target.id, student_id), select(ExamSubmission).where(ExamSubmission.exam_id == target.id,
                                                         ExamSubmission.student_id == student_id).order_by(ExamSubmission.id))
        if not rows:
            return source_result(0, {'kind': 'exam', 'id': target.id, 'confirmed_missing': True}) if missing_confirmed else pending('考试未参加，需要教师确认缺交', 'missing_pending')
        submission = rows[0]
        if submission.status != 'graded' or submission.score is None: return pending('考试尚未评分完成')
        full = sum((decimal(q.points) for q in fetch('exam_questions', target.id, select(ExamQuestion).where(ExamQuestion.exam_id == target.id).order_by(ExamQuestion.id))), Decimal(0))
        if full <= 0: return pending('考试满分必须大于0')
        return source_result(decimal(submission.score) / full * HUNDRED, {'kind': 'exam', 'id': target.id,
                             'submission_id': submission.id, 'score': str(submission.score), 'max_score': str(full),
                             'updated_at': str(submission.updated_at), 'graded_at': str(submission.graded_at)})
    query = select(ExperimentRecord).where(ExperimentRecord.offering_id == offering_id, ExperimentRecord.student_id == student_id)
    query = query.where(ExperimentRecord.lesson_id == target.entry_id) if target.entry_kind == 'lesson' else query.where(ExperimentRecord.module_id == target.entry_id)
    records = fetch('records', (target.entry_kind, target.entry_id, student_id), query.order_by(ExperimentRecord.id))
    if len(records) != 1:
        if not records and missing_confirmed: return source_result(0, {'kind': 'experiment', 'id': target.id, 'confirmed_missing': True})
        return pending('实验尚未提交或归属不唯一，需要教师确认', 'missing_pending')
    record = records[0]
    bindings = fetch('bindings', record.id, select(LabRecord).where(LabRecord.record_id == record.id))
    binding = bindings[0] if bindings else None
    if binding:
        # A previous selected score is not a final course result while a newer
        # formal submission is still grading or awaiting infrastructure review.
        states = fetch('states', record.id, select(LabSubmission).join(ExperimentSubmission,
            ExperimentSubmission.id == LabSubmission.submission_id)
            .where(ExperimentSubmission.record_id == record.id).order_by(LabSubmission.submission_id))
        grade_ids = [state.published_grade_id for state in states if state.published_grade_id]
        grade_rows = ([grade for ident in sorted(set(grade_ids)) for grade in source_rows['lab_grades'].get(ident, [])]
                      if source_rows is not None else query_rows(db, select(LabGradeRevision)
                          .where(LabGradeRevision.id.in_(grade_ids)).order_by(LabGradeRevision.id), lock)) if grade_ids else []
        grades = {grade.id: grade for grade in grade_rows}
        for state in states:
            published = grades.get(state.published_grade_id)
            if (state.review_requested or not published or not published.published_at
                    or published.submission_id != state.submission_id):
                runs = fetch('runs', state.current_run_id, select(LabRun).where(LabRun.id == state.current_run_id)) if state.current_run_id else []
                run = runs[0] if runs else None
                if run and run.status in ('queued', 'running', 'recovering'):
                    return pending('实验新提交仍在评分，完成后再确认总评')
                return pending('正式实验成绩尚未公开完成或基础设施失败，请完成复核或明确调分')
        if not binding.selected_submission_id:
            if missing_confirmed and not fetch('experiment_submissions', record.id, select(ExperimentSubmission).where(ExperimentSubmission.record_id == record.id).limit(1)):
                return source_result(0, {'kind': 'experiment', 'id': target.id, 'record_id': record.id, 'confirmed_missing': True})
            return pending('实验尚无平台选定成绩')
        selected = fetch('selected', binding.selected_submission_id, select(ExperimentSubmission).where(ExperimentSubmission.id == binding.selected_submission_id))
        if not selected or selected[0].record_id != record.id: return pending('实验选定记录无效')
        state = next((state for state in states if state.submission_id == selected[0].id), None)
        grade = grades.get(state.published_grade_id) if state else None
        if not grade or not grade.published_at or grade.submission_id != selected[0].id or state.review_requested:
            return pending('实验选定成绩尚未发布或等待复核')
        return source_result(grade.total, {'kind': 'experiment', 'id': target.id, 'record_id': record.id,
                             'submission_id': selected[0].id, 'grade_revision_id': grade.id,
                             'assessment_id': grade.assessment_id, 'selection': 'platform_selected',
                             'published_at': str(grade.published_at), 'score': str(grade.total)})
    submissions = fetch('experiment_submissions', record.id, select(ExperimentSubmission).where(ExperimentSubmission.record_id == record.id).order_by(ExperimentSubmission.id))
    if not submissions:
        return source_result(0, {'kind': 'experiment', 'id': target.id, 'confirmed_missing': True}) if missing_confirmed else pending('实验未提交，需要教师确认缺交', 'missing_pending')
    if any(s.score is None or s.reviewed_at is None for s in submissions): return pending('实验尚未完成教师评分')
    chosen = select_score(submissions, column.selection)
    return source_result(chosen.score, {'kind': 'experiment', 'id': target.id, 'record_id': record.id,
                         'submission_id': chosen.id, 'reviewed_at': str(chosen.reviewed_at), 'selection': column.selection})


def calculate(db, book, *, lock=False):
    groups, columns = categories(db, book, lock=lock)
    targets = {}
    for col in columns:
        if col.source_kind == 'manual': continue
        try: targets[col.id] = source_target(db, col.source_kind, col.source_id, book.offering_id, lock=lock)
        except HTTPException as exc:
            if exc.detail['code'] not in ('GRADE_SOURCE_SCOPE', 'GRADE_SOURCE_UNPUBLISHED'): raise
            targets[col.id] = None
    entries = {(e.column_id, e.student_id): e for e in query_rows(db, select(GradebookEntry).where(
        GradebookEntry.column_id.in_([c.id for c in columns])), lock)}
    adjustments = {a.student_id: a for a in query_rows(db, select(GradebookAdjustment).where(GradebookAdjustment.gradebook_id == book.id), lock)}
    rows, blockers = [], []
    students = roster(db, book.offering_id, lock=lock)
    source_rows = preload_sources(db, columns, targets, book.offering_id, [s.id for s in students], lock=lock)
    for student in students:
        cells, results, total, valid = {}, [], Decimal(0), True
        for group in groups:
            children = [c for c in columns if c.category_id == group.id]
            for col in children:
                entry = entries.get((col.id, student.id))
                if col.source_kind != 'manual' and targets.get(col.id) is None:
                    item = pending('原成绩来源不可用，请重新确认分项；官方历史版本保留')
                elif entry and entry.state == 'missing' and col.source_kind != 'manual':
                    item = source_score(db, col, book.offering_id, student.id, lock=lock, missing_confirmed=True, target=targets[col.id], source_rows=source_rows)
                    if item['source']: item['source']['missing_confirmation_reason'] = entry.reason
                elif entry and entry.state not in ('auto', 'pending'):
                    item = {'state': entry.state, 'score': score_text(entry.score) if entry.state == 'scored' else '0.00' if entry.state == 'missing' else None,
                            'message': entry.reason or '', 'source': {'kind': 'manual', 'entry_id': entry.id, 'updated_at': str(entry.updated_at), 'reason': entry.reason}}
                elif entry and entry.state == 'pending': item = pending('教师标记为待评分')
                elif col.source_kind == 'manual': item = pending('人工成绩待录入')
                else: item = source_score(db, col, book.offering_id, student.id, lock=lock, target=targets[col.id], source_rows=source_rows)
                cells[col.key] = item
            if group.weight == 0:
                results.append({'key': group.key, 'score': None, 'weighted_points': '0.00', 'effective_weights': {}})
                continue
            included = [c for c in children if cells[c.key]['state'] != 'exempt' and c.weight > 0]
            denominator = sum((c.weight for c in included), Decimal(0))
            if denominator == 0:
                valid = False; blockers.append({'student_id': student.id, 'category': group.key, 'message': '整个一级栏目免修，必须调整栏目配置后才能发布'})
                results.append({'key': group.key, 'score': None, 'weighted_points': None, 'effective_weights': {}})
                continue
            category_score, category_valid = Decimal(0), True
            weights = {c.key: str((c.weight / denominator * HUNDRED).quantize(Decimal('0.0001'))) for c in included}
            for col in included:
                cell = cells[col.key]
                if cell['score'] is None:
                    category_valid = valid = False
                    blockers.append({'student_id': student.id, 'column': col.key, 'message': cell['message']})
                else: category_score += decimal(cell['score']) * col.weight / denominator
            weighted = category_score * group.weight / HUNDRED if category_valid else None
            if weighted is not None: total += weighted
            results.append({'key': group.key, 'score': score_text(category_score) if category_valid else None,
                            'weighted_points': score_text(weighted), 'effective_weights': weights})
        adjustment = adjustments.get(student.id)
        rows.append({'student_id': student.id, 'student_no': student.student_no, 'student_name': student.real_name,
                     'cells': cells, 'categories': results, 'calculated_score': score_text(total) if valid else None,
                     'final_score': score_text(adjustment.score if adjustment else total) if valid else None,
                     'adjustment': {'score': score_text(adjustment.score), 'reason': adjustment.reason} if adjustment else None})
    result = {'offering_id': book.offering_id, 'book_revision': book.revision, 'config': config_read(db, book, lock=lock),
              'rows': rows, 'blockers': blockers}
    # 未决成绩异议只提示、不阻止发布（不进 blockers）：成绩册本身不提供异议入口，
    # 教师看到计数后到 AI 评分复核页处理。
    result['warnings'] = _appeal_warnings(db, book.offering_id)
    reviewed = {key: value for key, value in result.items() if key != 'book_revision'}
    result['preview_hash'] = hashlib.sha256(json.dumps(reviewed, sort_keys=True, ensure_ascii=False, separators=(',', ':')).encode()).hexdigest()
    return result


def _appeal_warnings(db, offering_id: int) -> list[dict]:
    from app.services.appeal_service import count_offering_pending_appeals

    pending = count_offering_pending_appeals(db, offering_id)
    if not pending:
        return []
    return [{
        'kind': 'PENDING_APPEALS',
        'count': pending,
        'message': f'有 {pending} 条成绩异议尚未处理，不影响总评发布，建议发布前先在 AI 评分复核页处理',
    }]


def snapshot_read(db, snapshot, actor):
    book = db.get(Gradebook, snapshot.gradebook_id)
    payload = deepcopy(snapshot.payload)
    if actor.role == 'student':
        payload['rows'] = [r for r in payload['rows'] if r['student_id'] == actor.id]
        if not payload['rows']: raise api_error(403, 'FORBIDDEN', '只能查看自己的已发布成绩')
        payload.pop('blockers', None)
    else: require_offering(db, book.offering_id, actor)
    return {'id': snapshot.id, 'revision': snapshot.revision, 'kind': snapshot.kind, 'created_at': snapshot.created_at,
            'reason': snapshot.reason, 'payload_hash': snapshot.payload_hash, 'source_snapshot_id': snapshot.source_snapshot_id, 'payload': payload}


def publish(db, offering_id, actor, payload):
    book = lock_book(db, offering_id, actor, payload.expected_revision)
    result = calculate(db, book, lock=True)
    if not payload.expected_preview_hash:
        raise api_error(422, 'GRADEBOOK_PREVIEW_REQUIRED', '请先加载并审核当前成绩草稿再发布')
    if result['preview_hash'] != payload.expected_preview_hash:
        raise api_error(409, 'GRADEBOOK_PREVIEW_CHANGED', '学生名单或来源成绩已变化，请刷新草稿并重新审核')
    if not result['rows']: raise api_error(409, 'GRADEBOOK_EMPTY_ROSTER', '授课实例暂无学生，不能发布成绩')
    if result['blockers']:
        raise api_error(409, 'GRADEBOOK_INCOMPLETE', '存在待评分、缺交未确认或整栏免修，不能发布', fields={'blockers': result['blockers'][:100]})
    digest = hashlib.sha256(json.dumps(result, sort_keys=True, ensure_ascii=False, separators=(',', ':')).encode()).hexdigest()
    snapshot = GradebookSnapshot(id=str(uuid4()), gradebook_id=book.id, revision=book.revision,
        kind='published', payload=result, payload_hash=digest, reason=payload.reason, actor_id=actor.id)
    db.add(snapshot); book.published_snapshot_id = snapshot.id; db.flush()
    audit(db, actor.id, 'gradebook.publish', 'gradebook', book.id, reason=payload.reason,
          changes={'snapshot_id': snapshot.id, 'revision': book.revision, 'count': len(result['rows']), 'payload_hash': digest})
    return snapshot


def archive(db, offering_id, actor, payload):
    book = lock_book(db, offering_id, actor, payload.expected_revision, write=False)
    sources = query_rows(db, select(GradebookSnapshot).where(GradebookSnapshot.id == book.published_snapshot_id), True) if book.published_snapshot_id else []
    source = sources[0] if sources else None
    if not source: raise api_error(409, 'GRADEBOOK_NOT_PUBLISHED', '请先发布一个完整成绩版本，再进行归档')
    if source.payload['book_revision'] != payload.expected_revision:
        last = db.scalar(select(GradebookSnapshot).where(GradebookSnapshot.gradebook_id == book.id)
                         .order_by(GradebookSnapshot.revision.desc()).limit(1).with_for_update()
                         .execution_options(populate_existing=True))
        if last and last.kind == 'archive' and last.source_snapshot_id == source.id and last.revision == payload.expected_revision:
            raise api_error(409, 'GRADEBOOK_ALREADY_ARCHIVED', '该官方成绩版本已归档，无需重复归档')
        # Archive only the latest publication; unpublished edits cannot be silently dropped.
        raise api_error(409, 'GRADEBOOK_UNPUBLISHED_CHANGES', '当前有未发布修改，请先发布或明确恢复原配置后再归档')
    snapshot = GradebookSnapshot(id=str(uuid4()), gradebook_id=book.id, revision=book.revision,
        kind='archive', payload=deepcopy(source.payload), payload_hash=source.payload_hash,
        source_snapshot_id=source.id, reason=payload.reason, actor_id=actor.id)
    db.add(snapshot); db.flush()
    audit(db, actor.id, 'gradebook.archive', 'gradebook', book.id, reason=payload.reason,
          changes={'snapshot_id': snapshot.id, 'source_snapshot_id': source.id, 'revision': book.revision})
    return snapshot


def source_options(db, offering_id):
    options = []
    for kind, model in [('assignment', Assignment), ('exam', Exam)]:
        options.extend({'kind': kind, 'id': r.id, 'title': r.title} for r in db.scalars(select(model).where(model.offering_id == offering_id, model.status == 'published')))
    for r in db.scalars(select(OfferingExperiment).where(OfferingExperiment.offering_id == offering_id, OfferingExperiment.status == 'published')):
        entry = db.get(Lesson if r.entry_kind == 'lesson' else ExperimentModule, r.entry_id)
        if entry: options.append({'kind': 'experiment', 'id': r.id, 'title': entry.title if r.entry_kind == 'lesson' else entry.name})
    return options


def read_book(db, offering_id, actor):
    book = book_for(db, offering_id)
    latest = db.get(GradebookSnapshot, book.published_snapshot_id) if book and book.published_snapshot_id else None
    if actor.role == 'student':
        if book and (not latest or not any(r['student_id'] == actor.id for r in latest.payload['rows'])):
            latest = next((snapshot for snapshot in db.scalars(select(GradebookSnapshot).where(
                GradebookSnapshot.gradebook_id == book.id, GradebookSnapshot.kind == 'published')
                .order_by(GradebookSnapshot.revision.desc()))
                if any(r['student_id'] == actor.id for r in snapshot.payload['rows'])), None)
        if not latest: require_offering(db, offering_id, actor)
        return {'offering_id': offering_id, 'snapshot': snapshot_read(db, latest, actor) if latest else None}
    offering = require_offering(db, offering_id, actor)
    term, klass = db.get(AcademicTerm, offering.academic_term_id), db.get(TeachingClass, offering.teaching_class_id)
    return {'offering_id': offering_id, 'revision': book.revision if book else 0,
            'read_only': offering.status != 'active' or not term or term.status == 'closed' or not klass or klass.status == 'archived',
            'config': config_read(db, book) if book else default_config(),
            'draft': calculate(db, book) if book else None,
            'snapshot': snapshot_read(db, latest, actor) if latest else None,
            'sources': source_options(db, offering_id)}


def export_snapshot(db, snapshot, actor, format):
    data = snapshot_read(db, snapshot, actor)
    config = data['payload']['config']
    columns = [c for group in config for c in group['columns']]
    headers = ['学号', '姓名', *[c['title'] for c in columns], '计算总分', '最终总分', '调分原因', '成绩版本']
    values = [[row['student_no'] or '', row['student_name'], *[
        '免修' if row['cells'][c['key']]['state'] == 'exempt' else row['cells'][c['key']]['score'] for c in columns],
        row['calculated_score'], row['final_score'], row['adjustment']['reason'] if row['adjustment'] else '', str(data['revision'])]
        for row in data['payload']['rows']]
    audit(db, actor.id, 'gradebook.export', 'gradebook_snapshot', snapshot.id, changes={'count': len(values), 'format': format})
    if format == 'csv':
        stream = StringIO(); writer = csv.writer(stream)
        def safe(value):
            text = str(value if value is not None else '')
            return "'" + text if text.lstrip().startswith(('=', '+', '-', '@')) else text
        writer.writerow([safe(v) for v in headers]); writer.writerows([[safe(v) for v in row] for row in values])
        return ('\ufeff' + stream.getvalue()).encode('utf-8')
    book = Workbook(); sheet = book.active; sheet.title = '课程成绩'
    sheet.append(headers)
    for cell in sheet[1]: cell.data_type = 's'; cell.number_format = '@'
    for values_row in values:
        sheet.append([str(v if v is not None else '') for v in values_row])
        for cell in sheet[sheet.max_row]: cell.data_type = 's'; cell.number_format = '@'
    sheet.freeze_panes = 'C2'
    from openpyxl.utils import get_column_letter
    for col in range(1, len(headers) + 1): sheet.column_dimensions[get_column_letter(col)].width = 24
    out = BytesIO(); book.save(out)
    return out.getvalue()
