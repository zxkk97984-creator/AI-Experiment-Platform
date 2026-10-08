"""Credential issuance and metadata projection. Callers own transactions."""
import secrets
from datetime import timedelta
from io import BytesIO
from uuid import uuid4

from openpyxl import Workbook
from sqlalchemy import select, update

from app.errors import api_error
from app.models import User
from app.models.credential_delivery import CredentialDelivery
from app.services.time_utils import as_utc, utc_now

PASSWORD_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZabcdefghjkmnpqrstuvwxyz23456789"


def generate_password(identity=None, used=None):
    used = used if used is not None else set()
    while True:
        password = ''.join(secrets.choice(PASSWORD_ALPHABET) for _ in range(16))
        if password not in used and password.casefold() != (identity or '').casefold():
            used.add(password)
            return password


def prepare(db, actor_id, settings, recipient_key, *, batch_id=None, used=None, identity=None):
    from app.services.account_service import encrypt_secret, decrypt_secret
    delivery = db.scalar(select(CredentialDelivery).where(
        CredentialDelivery.batch_id == batch_id, CredentialDelivery.recipient_key == recipient_key)) if batch_id else None
    if delivery:
        if not delivery.ciphertext:
            raise api_error(409, 'CREDENTIAL_UNAVAILABLE', '临时凭据已清理，请重新发起操作')
        return delivery, decrypt_secret(settings, delivery.ciphertext, f'credential:{delivery.id}')
    delivery = CredentialDelivery(id=str(uuid4()), actor_id=actor_id, batch_id=batch_id,
                                  recipient_key=recipient_key, state='prepared')
    password = generate_password(identity or recipient_key.split(':', 1)[-1], used)
    delivery.ciphertext = encrypt_secret(settings, password, f'credential:{delivery.id}')
    db.add(delivery)
    db.flush()
    return delivery, password


def issue(db, delivery, user, settings):
    invalidate(db, user.id, except_id=delivery.id)
    delivery.user_id = user.id
    delivery.state = 'issued'
    delivery.issued_credential_version = user.credential_version
    delivery.issued_at = utc_now()
    delivery.export_expires_at = delivery.issued_at + timedelta(hours=getattr(settings, 'account_credential_export_hours', 24))
    db.flush()


def invalidate(db, user_id, except_id=None):
    filters = [CredentialDelivery.user_id == user_id, CredentialDelivery.ciphertext.is_not(None)]
    if except_id:
        filters.append(CredentialDelivery.id != except_id)
    db.execute(update(CredentialDelivery).where(*filters).values(ciphertext=None, state='cleared'))


def eligible(delivery, user, now=None):
    now = now or utc_now()
    return bool(delivery.state == 'issued' and delivery.ciphertext and delivery.export_expires_at
                and as_utc(delivery.export_expires_at) > now and user and user.role == 'student'
                and user.status == 'active' and user.must_change_password
                and user.credential_version == delivery.issued_credential_version
                and user.temporary_password_expires_at and as_utc(user.temporary_password_expires_at) > now)


def metadata(db, *, batch_id=None, user_id=None):
    query = select(CredentialDelivery, User).outerjoin(User, User.id == CredentialDelivery.user_id)
    query = query.where(CredentialDelivery.batch_id == batch_id) if batch_id else query.where(CredentialDelivery.user_id == user_id)
    return [{'id': d.id, 'user_id': d.user_id, 'issued_at': as_utc(d.issued_at),
             'export_expires_at': as_utc(d.export_expires_at), 'available': eligible(d, u)}
            for d, u in db.execute(query.order_by(CredentialDelivery.issued_at.desc(), CredentialDelivery.created_at.desc(), CredentialDelivery.id)).all()
            if d.state != 'prepared']


def workbook(db, rows, actor, settings, *, target_type, target_id):
    from app.services.account_service import audit, decrypt_secret
    current = [(d, u) for d, u in rows if eligible(d, u)]
    if not current:
        raise api_error(410, 'CREDENTIAL_DELIVERY_EXPIRED', '没有可导出的临时凭据，可能已过期、改密或被重新重置')
    book = Workbook()
    sheet = book.active
    sheet.title = '临时凭据'
    sheet.append(['学号', '姓名', '登录账号', '独立临时密码', '首次改密截止时间'])
    for delivery, user in current:
        values = [user.student_no or '', user.real_name, user.username,
                  decrypt_secret(settings, delivery.ciphertext, f'credential:{delivery.id}'),
                  as_utc(user.temporary_password_expires_at).isoformat()]
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'
            cell.number_format = '@'
    sheet.freeze_panes = 'A2'
    from openpyxl.utils import get_column_letter
    for col in range(1, 6):
        sheet.column_dimensions[get_column_letter(col)].width = 28
    output = BytesIO()
    book.save(output)
    audit(db, actor.id, 'credential.export', target_type, target_id,
          changes={'count': len(current)})
    return output.getvalue()


def cleanup(db):
    now = utc_now()
    db.execute(update(CredentialDelivery).where(
        CredentialDelivery.state == 'issued', CredentialDelivery.export_expires_at <= now
    ).values(ciphertext=None, state='cleared'))
    current = select(User.id).where(User.id == CredentialDelivery.user_id, User.role == 'student',
                                  User.status == 'active', User.must_change_password.is_(True),
                                  User.credential_version == CredentialDelivery.issued_credential_version,
                                  User.temporary_password_expires_at > now).exists()
    db.execute(update(CredentialDelivery).where(CredentialDelivery.state == 'issued', ~current)
               .values(ciphertext=None, state='cleared'))
