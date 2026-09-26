"""Strict XLSX parsing and deterministic, side-effect-free import preflight."""
from collections import Counter
from io import BytesIO
from zipfile import BadZipFile, ZipFile
import hashlib
import json

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.errors import api_error
from app.models import AcademicTerm, TeachingClass, TeachingClassStudent, CourseTeachingClass, User

HEADERS = {'学号': 'student_no', '姓名': 'real_name', '院系': 'department',
           '班级编码': 'class_code', '班级名称': 'class_name'}
MAX_ROWS = 5000


def spreadsheet(kind, rows=None):
    headers = ['班级编码', '班级名称'] if kind == 'classes' else ['学号', '姓名', '院系', '班级编码']
    if kind == 'class_students':
        headers = headers[:3]
    if rows is not None:
        headers = ['行号', '状态', '说明', '学号', '姓名', '班级编码']
    wb = Workbook()
    ws = wb.active
    ws.title = '名单' if rows is None else '结果'
    ws.append(headers)
    if rows is None:
        for col in range(1, len(headers) + 1):
            for row in range(2, MAX_ROWS + 2):
                ws.cell(row, col).number_format = '@'
    else:
        for values in rows:
            for col, value in enumerate(values, start=1):
                cell = ws.cell(ws.max_row + 1 if col == 1 else ws.max_row, col)
                cell.value = str(value if value is not None else '')
                cell.data_type = 's'
    ws.freeze_panes = 'A2'
    from openpyxl.utils import get_column_letter
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 24
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def parse_xlsx(content, filename, kind):
    if not filename or not filename.lower().endswith('.xlsx') or len(content) > 5 * 1024 * 1024:
        raise api_error(422, 'INVALID_XLSX', '仅支持不超过 5 MiB 的 .xlsx 文件')
    try:
        from defusedxml.ElementTree import fromstring
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > 50 * 1024 * 1024:
                raise ValueError('工作簿解压后过大')
            if any('vbaproject' in e.filename.lower() or 'externallinks/' in e.filename.lower() for e in entries):
                raise ValueError('不支持宏或外部链接')
            for entry in entries:
                if entry.filename.endswith('.rels'):
                    if any(node.attrib.get('TargetMode') == 'External' for node in fromstring(archive.read(entry))):
                        raise ValueError('不支持外部链接')
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if len(wb.worksheets) != 1:
                raise ValueError('只允许一个工作表')
            ws = wb.worksheets[0]
            ws.reset_dimensions()  # Do not trust producer-supplied worksheet dimensions.
            iterator = ws.iter_rows()
            first = next(iterator, ())
            if len(first) > 20:
                raise ValueError('最多允许 20 列')
            headers = [str(c.value or '').strip() for c in first]
            while headers and not headers[-1]:
                headers.pop()
            if len(headers) != len(set(headers)) or any(h not in HEADERS for h in headers):
                raise ValueError('请使用模板表头，不允许重复列或额外列')
            required = {'班级编码', '班级名称'} if kind == 'classes' else {'学号', '姓名'}
            if not required.issubset(headers):
                raise ValueError('缺少必填表头')
            output = []
            scanned = 0
            for number, cells in enumerate(iterator, 2):
                scanned += 1
                if scanned > MAX_ROWS:
                    raise ValueError('最多允许 5000 行数据')
                if len(cells) > 20:
                    raise ValueError('最多允许 20 列')
                if not any(c.value is not None for c in cells):
                    continue
                data = {}
                errors = []
                for col, cell in enumerate(cells):
                    if cell.value is None:
                        continue
                    if col >= len(headers):
                        errors.append('存在无表头数据')
                        continue
                    key = HEADERS[headers[col]]
                    if cell.data_type == 'f':
                        errors.append('不允许公式')
                    if key in {'student_no', 'class_code'} and not isinstance(cell.value, str):
                        errors.append(f'{headers[col]}必须为文本，不能恢复已丢失的前导零')
                    value = str(cell.value).strip()
                    if len(value) > (160 if key == 'class_name' else 120 if key in {'real_name', 'department'} else 80):
                        errors.append(f'{headers[col]}超长')
                    data[key] = value
                data['_parse_errors'] = errors
                output.append((number, data))
            if not output:
                raise ValueError('工作簿没有数据')
            return output
        finally:
            wb.close()
    except (BadZipFile, ValueError, KeyError, OSError, TypeError) as exc:
        raise api_error(422, 'INVALID_XLSX', str(exc))
    except Exception:
        # Parser internals and XML paths must not be returned to the client.
        raise api_error(422, 'INVALID_XLSX', '工作簿损坏或包含不支持的 XML 数据')


def preflight(db, kind, scope, inputs, *, locking=False):
    def query(statement):
        return statement.with_for_update().execution_options(populate_existing=True) if locking else statement
    def one(model, ident):
        return db.scalar(query(select(model).where(model.id == ident))) if ident else None
    term_id, class_id = scope.get('term_id'), scope.get('class_id')
    target = one(TeachingClass, class_id)
    if class_id and not target:
        raise api_error(422, 'INVALID_CLASS', '教学班不存在')
    if target:
        term_id = target.academic_term_id
    term = one(AcademicTerm, term_id)
    if (term_id and not term) or (kind == 'classes' and not term):
        raise api_error(422, 'INVALID_TERM', '请选择有效学期')
    if term and term.status == 'closed':
        raise api_error(409, 'ACADEMIC_TERM_CLOSED', '学期教务名单已冻结')
    if target and target.status != 'active':
        raise api_error(409, 'TEACHING_CLASS_ARCHIVED', '教学班已归档')
    classes = {c.code: c for c in db.scalars(query(select(TeachingClass).where(
        TeachingClass.academic_term_id == term_id))).all()} if term else {}
    numbers = {d.get('student_no') for _, d in inputs if d.get('student_no')}
    users = {}
    usernames = {}
    # Chunk IN queries to work with SQLite limits as well as MySQL.
    number_list = list(numbers)
    for start in range(0, len(number_list), 500):
        part = number_list[start:start + 500]
        users.update({u.student_no: u for u in db.scalars(query(select(User).where(User.student_no.in_(part)))).all()})
        usernames.update({u.username: u for u in db.scalars(query(select(User).where(User.username.in_(part)))).all()})
    member_map = {(m.teaching_class_id, m.student_id): m.status for m in db.scalars(
        query(select(TeachingClassStudent).where(TeachingClassStudent.teaching_class_id.in_([c.id for c in classes.values()])))).all()} if classes else {}
    seen = set()
    identities = {}
    totals = Counter()
    unique_users = set()
    results = []
    for number, data in inputs:
        errors = list(data.get('_parse_errors', []))
        code = target.code if target else data.get('class_code', '')
        cls = classes.get(code)
        result = {'row': number, 'class_id': cls.id if cls else None, 'class_code': code}
        if kind == 'classes':
            name = data.get('class_name', '')
            if not code or not name:
                errors.append('班级编码和名称必填')
            if cls and (cls.name != name or cls.status != 'active'):
                errors.append('班级名称冲突或已归档')
            if code in identities and identities[code] != name:
                errors.append('同编码班级名称不一致')
            identities[code] = name
            result.update(action='unchanged' if cls else 'create_class', observed=[cls.name, cls.status] if cls else None)
            duplicate = code in seen
            seen.add(code)
        else:
            no, name, department = data.get('student_no', ''), data.get('real_name', ''), data.get('department', '')
            user = users.get(no)
            if not no or not name:
                errors.append('学号和姓名必填')
            identity = identities.get(no)
            if identity and (identity[0] != name or (identity[1] and department and identity[1] != department)):
                errors.append('同学号基础资料不一致')
            identities[no] = (name, department or (identity[1] if identity else ''))
            if user:
                if user.role != 'student' or user.status != 'active':
                    errors.append('账号不是有效学生')
                if user.real_name != name or (department and user.department and department != user.department):
                    errors.append('已有学生姓名或院系冲突')
            elif no in usernames:
                errors.append('学号对应的登录名已被占用')
            if code and (not cls or cls.status != 'active'):
                errors.append('班级不存在或已归档，请先导入班级')
            if target and data.get('class_code') and data['class_code'] != target.code:
                errors.append('文件班级与当前班级不一致')
            member = member_map.get((cls.id, user.id)) if cls and user else None
            result.update(user_id=user.id if user else None,
                          action='reuse_user' if user else 'create_user',
                          membership='unchanged' if member == 'active' else 'restore' if member else 'add' if cls else None,
                          fill_department=bool(user and not user.department and department),
                          observed=[user.id, user.real_name, user.department, user.role, user.status, member] if user else None)
            duplicate = (no, code) in seen
            seen.add((no, code))
            if not errors and no not in unique_users:
                totals['reused_accounts' if user else 'new_accounts'] += 1
                unique_users.add(no)
        result['errors'] = errors
        if duplicate:
            result['duplicate'] = True
        if errors:
            totals['errors'] += 1
        elif duplicate:
            totals['unchanged'] += 1
        else:
            totals[result['action']] += 1
            if result.get('membership') in {'add', 'restore'}:
                totals['memberships'] += 1
            if result.get('fill_department'):
                totals['filled_departments'] += 1
        results.append(result)
    links = list(db.execute(query(select(CourseTeachingClass.course_id, CourseTeachingClass.teaching_class_id)
                            .where(CourseTeachingClass.teaching_class_id.in_([c.id for c in classes.values()]))
                            .order_by(CourseTeachingClass.course_id, CourseTeachingClass.teaching_class_id))).all()) if classes else []
    digest = hashlib.sha256(json.dumps({'results': results, 'links': [list(x) for x in links],
                                      'term': [term.id, term.status] if term else None},
                                     sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return results, dict(totals), digest
