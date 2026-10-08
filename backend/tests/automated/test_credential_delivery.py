"""Real issuance, retry, export and revocation across API/worker/database boundaries."""
from datetime import timedelta
from io import BytesIO

from cryptography.fernet import Fernet
from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import func, select

from app.models import AdminBatch, User
from app.models.credential_delivery import CredentialDelivery
from app.security import verify_password
from app.services import account_batch_service as worker
from app.services import account_service as accounts
from app.services import credential_delivery_service as credentials
from app.services.time_utils import as_utc, utc_now
from conftest import auth_header, create_user, login

API = '/api/v1'


@pytest.fixture()
def admin(client, db_session_factory, test_settings):
    test_settings.account_security_keys = Fernet.generate_key().decode()
    user = create_user(db_session_factory, 'credential-admin', 'admin')
    return user, auth_header(login(client, user.username)[0])


def read_sheet(response):
    assert response.status_code == 200, response.text
    assert response.headers['cache-control'] == 'private, no-store'
    book = load_workbook(BytesIO(response.content), read_only=True)
    try:
        return list(book.active.iter_rows(min_row=2, values_only=True))
    finally:
        book.close()


def create_student(client, headers, no='CD001', name='Student'):
    response = client.post(f'{API}/users', headers=headers,
                           json={'student_no': no, 'real_name': name, 'role': 'student'})
    assert response.status_code == 201, response.text
    assert 'password' not in response.json() and 'ciphertext' not in response.json()
    return response.json()


def delivery(client, headers, user_id):
    data = client.get(f'{API}/admin/credential-deliveries', headers=headers, params={'user_id': user_id})
    assert data.status_code == 200
    assert 'ciphertext' not in data.text
    return data.json()['items'][0]


def export(client, headers, ident, password='Passw0rd!'):
    return client.post(f'{API}/admin/credential-deliveries/{ident}/export', headers=headers,
                       json={'current_password': password})


def import_batch(client, headers, rows):
    book = Workbook(); book.active.append(['学号', '姓名'])
    for row in rows: book.active.append(row)
    out = BytesIO(); book.save(out)
    response = client.post(f'{API}/admin/roster-imports/preview', headers=headers,
                           data={'kind': 'students'}, files={'file': ('students.xlsx', out.getvalue())})
    assert response.status_code == 200, response.text
    batch = response.json()
    confirmed = client.post(f"{API}/admin/batches/{batch['id']}/confirm", headers=headers,
                            json={'preview_digest': batch['preview_digest']})
    assert confirmed.status_code == 202, confirmed.text
    return batch


def test_single_issuance_is_unique_exportable_and_requires_first_password_change(client, db_session_factory, admin):
    _, headers = admin
    first = create_student(client, headers, name='=1+1')
    second = create_student(client, headers, no='CD002')
    first_meta = delivery(client, headers, first['id'])
    second_meta = delivery(client, headers, second['id'])
    first_export = export(client, headers, first_meta['id'])
    first_password = read_sheet(first_export)[0][3]
    second_password = read_sheet(export(client, headers, second_meta['id']))[0][3]
    assert first_password != second_password
    assert len(first_password) == len(second_password) == 16
    with db_session_factory() as db:
        user = db.get(User, first['id'])
        record = db.get(CredentialDelivery, first_meta['id'])
        assert verify_password(first_password, user.password_hash)
        assert as_utc(record.export_expires_at) - as_utc(record.issued_at) == timedelta(hours=24)
        assert timedelta(days=7) - timedelta(seconds=2) < as_utc(user.temporary_password_expires_at) - as_utc(record.issued_at) <= timedelta(days=7)
    book = load_workbook(BytesIO(first_export.content))
    assert book.active['B2'].data_type == 's'  # User text cannot become an Excel formula.
    token, _ = login(client, first['username'], first_password)
    assert client.get(f'{API}/courses', headers=auth_header(token)).status_code == 403
    changed = client.patch(f'{API}/users/me/password', headers=auth_header(token),
                           json={'current_password': first_password, 'password': 'StudentPersonal!123'})
    assert changed.status_code == 204
    assert export(client, headers, first_meta['id']).status_code == 410
    with db_session_factory() as db:
        assert db.get(CredentialDelivery, first_meta['id']).ciphertext is None
    assert client.get(f'{API}/auth/me', headers=auth_header(token)).status_code == 401


def test_export_requires_admin_reauthentication_and_rejects_expiry(client, db_session_factory, admin):
    _, headers = admin
    student = create_student(client, headers)
    meta = delivery(client, headers, student['id'])
    assert export(client, headers, meta['id'], 'IncorrectPassword!').status_code == 422
    teacher = create_user(db_session_factory, 'credential-teacher', 'teacher')
    assert export(client, auth_header(login(client, teacher.username)[0]), meta['id']).status_code == 403
    with db_session_factory() as db:
        db.get(CredentialDelivery, meta['id']).export_expires_at = utc_now() - timedelta(seconds=1)
        db.commit()
    assert export(client, headers, meta['id']).status_code == 410
    assert delivery(client, headers, student['id'])['available'] is False
    with db_session_factory() as db:
        worker.cleanup_batches(db, type('Retention', (), {'account_batch_retention_days': 90, 'account_audit_retention_days': 365})())
        assert db.get(CredentialDelivery, meta['id']).ciphertext is None


def test_reset_revokes_prior_delivery_and_sessions_and_issues_fresh_password(client, db_session_factory, admin):
    _, headers = admin
    student = create_student(client, headers)
    old = delivery(client, headers, student['id'])
    old_password = read_sheet(export(client, headers, old['id']))[0][3]
    token, _ = login(client, student['username'], old_password)
    result = client.post(f"{API}/users/{student['id']}/password-reset", headers=headers,
                         json={'current_password': 'Passw0rd!', 'reason': 'Student requested reset'})
    assert result.status_code == 200, result.text
    new = delivery(client, headers, student['id'])
    new_password = read_sheet(export(client, headers, new['id']))[0][3]
    assert new_password != old_password
    assert len(new_password) == 16
    assert export(client, headers, old['id']).status_code == 410
    assert client.get(f'{API}/auth/me', headers=auth_header(token)).status_code == 401
    assert login(client, student['username'], new_password)


def test_batch_recovery_reuses_credentials_and_duplicate_rows_deliver_once(client, db_session_factory, test_settings, admin):
    _, headers = admin
    batch = import_batch(client, headers, [['BR001', 'A'], ['BR002', 'B'], ['BR001', 'A']])
    with db_session_factory() as db: first_claim = worker.claim_batch(db)
    worker.prepare_hashes(db_session_factory, *first_claim, test_settings)
    with db_session_factory() as db:
        prepared = {d.recipient_key: accounts.decrypt_secret(test_settings, d.ciphertext, f'credential:{d.id}')
                    for d in db.scalars(select(CredentialDelivery).where(CredentialDelivery.batch_id == batch['id']))}
        assert len(prepared) == 2 and len(set(prepared.values())) == 2
        assert all(len(password) == 16 for password in prepared.values())
        db.get(AdminBatch, batch['id']).lease_until = utc_now() - timedelta(minutes=3)
        db.commit()
    with db_session_factory() as db: new_claim = worker.claim_batch(db)
    worker.process_batch(db_session_factory, *first_claim, test_settings)
    worker.process_batch(db_session_factory, *new_claim, test_settings)
    rows = read_sheet(client.post(f"{API}/admin/batches/{batch['id']}/credentials", headers=headers,
                                 json={'current_password': 'Passw0rd!'}))
    assert {f'student:{row[0]}': row[3] for row in rows} == prepared
    ordinary = client.get(f"{API}/admin/batches/{batch['id']}", headers=headers)
    assert all(password not in ordinary.text for password in prepared.values())
    with db_session_factory() as db:
        assert db.scalar(select(func.count()).select_from(User).where(User.role == 'student')) == 2
        for row in rows:
            assert verify_password(row[3], db.scalar(select(User).where(User.student_no == row[0])).password_hash)


def test_missing_key_cannot_create_student_or_plaintext_delivery(client, db_session_factory, test_settings, admin):
    _, headers = admin
    test_settings.account_security_keys = ''
    result = client.post(f'{API}/users', headers=headers, json={'student_no': 'NO_KEY', 'real_name': 'A', 'role': 'student'})
    assert result.status_code == 503
    with db_session_factory() as db:
        assert db.scalar(select(func.count()).select_from(CredentialDelivery)) == 0
        assert db.scalar(select(User).where(User.student_no == 'NO_KEY')) is None


def test_batch_reset_exports_only_success_and_never_replays_successful_credentials(client, db_session_factory, test_settings, admin):
    _, headers = admin
    one = create_user(db_session_factory, 'reset-one', 'student')
    two = create_user(db_session_factory, 'reset-two', 'student')
    preview = client.post(f'{API}/admin/password-resets/preview', headers=headers,
                          json={'student_ids': [one.id, two.id]}).json()
    assert client.post(f"{API}/admin/batches/{preview['id']}/confirm", headers=headers,
                       json={'preview_digest': preview['preview_digest'], 'current_password': 'Passw0rd!', 'reason': 'New semester'}).status_code == 202
    with db_session_factory() as db:
        claim = worker.claim_batch(db)
        db.get(User, two.id).status = 'disabled'
        db.commit()
    worker.process_batch(db_session_factory, *claim, test_settings)
    response = client.post(f"{API}/admin/batches/{preview['id']}/credentials", headers=headers,
                           json={'current_password': 'Passw0rd!'})
    rows = read_sheet(response)
    assert len(rows) == 1 and rows[0][2] == one.username
    with db_session_factory() as db:
        assert verify_password(rows[0][3], db.get(User, one.id).password_hash)
        assert db.get(AdminBatch, preview['id']).status == 'partially_succeeded'
        skipped = db.scalar(select(CredentialDelivery).where(CredentialDelivery.batch_id == preview['id'],
                             CredentialDelivery.recipient_key == f'user:{two.id}'))
        assert skipped.ciphertext is None
        original_hash = db.get(User, one.id).password_hash
    worker.process_batch(db_session_factory, *claim, test_settings)
    with db_session_factory() as db:
        assert db.get(User, one.id).password_hash == original_hash
    audit = client.get(f'{API}/admin/audit-events', headers=headers, params={'action': 'credential.export'}).json()['items']
    assert len(audit) == 1 and audit[0]['changes'] == {'count': 1}
    assert rows[0][3] not in str(audit)


def test_random_temporary_password_rejects_batch_collision(monkeypatch):
    choices = iter('A' * 16 + 'B' * 16)
    monkeypatch.setattr(credentials.secrets, 'choice', lambda _alphabet: next(choices))
    used = {'A' * 16}
    password = credentials.generate_password('000123', used)
    assert password == 'B' * 16
    assert used == {'A' * 16, 'B' * 16}
