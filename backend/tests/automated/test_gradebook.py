from copy import deepcopy
import csv
from datetime import date
from io import BytesIO, StringIO
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest
from sqlalchemy import select

from app.api import gradebook as api
from app.models import (AcademicTerm, Assignment, Course, Exam, ExamQuestion, ExamSubmission,
                        JudgeQuestion, Submission, TeachingClass, User)
from app.models.gradebook import Gradebook, GradebookSnapshot
from app.models.teaching import CourseOffering, OfferingEnrollment
from app.services.gradebook_service import default_config
from conftest import auth_header, create_user, login

API = '/api/v1/gradebooks'


@pytest.fixture()
def case(app, client, db_session_factory):
    if not any(getattr(r, 'path', None) == f'{API}/offerings' for r in app.routes):
        app.include_router(api.router, prefix='/api/v1')
    teacher = create_user(db_session_factory, 'book-teacher', 'teacher')
    other = create_user(db_session_factory, 'book-other', 'teacher')
    owner = create_user(db_session_factory, 'book-owner', 'teacher')
    student = create_user(db_session_factory, 'book-student', 'student', real_name='=Student')
    second = create_user(db_session_factory, 'book-second', 'student', real_name='Other Student Private')
    with db_session_factory() as db:
        db.get(User, student.id).student_no = '00001'
        db.get(User, second.id).student_no = '00002'
        term = AcademicTerm(code='BOOK-TERM', name='Term', start_date=date(2026, 9, 1), end_date=date(2027, 1, 1), status='active')
        db.add(term); db.flush()
        course = Course(title='Shared course', teacher_id=owner.id, status='published')
        db.add(course); db.flush()
        klass = TeachingClass(academic_term_id=term.id, code='BOOK-1', name='Class one', status='active')
        alternate = TeachingClass(academic_term_id=term.id, code='BOOK-2', name='Class two', status='active')
        db.add_all([klass, alternate]); db.flush()
        offering = CourseOffering(course_id=course.id, academic_term_id=term.id, teaching_class_id=klass.id, teacher_id=teacher.id, status='active')
        other_offering = CourseOffering(course_id=course.id, academic_term_id=term.id, teaching_class_id=alternate.id, teacher_id=other.id, status='active')
        db.add_all([offering, other_offering]); db.flush()
        for user in (student, second):
            db.add(OfferingEnrollment(offering_id=offering.id, student_id=user.id, status='active'))
        db.commit()
        ident, other_id, course_id = offering.id, other_offering.id, course.id
    return SimpleNamespace(teacher=teacher, other=other, owner=owner, student=student, second=second,
        offering=ident, other_offering=other_id, course=course_id,
        headers=auth_header(login(client, teacher.username)[0]),
        student_headers=auth_header(login(client, student.username)[0]))


def save(client, case, config=None, revision=0):
    response = client.put(f'{API}/offerings/{case.offering}/config', headers=case.headers,
                          json={'expected_revision': revision, 'categories': config or default_config()})
    assert response.status_code == 200, response.text
    return response.json()


def edit(client, case, entries, revision):
    response = client.patch(f'{API}/offerings/{case.offering}/entries', headers=case.headers,
                            json={'expected_revision': revision, 'entries': entries})
    assert response.status_code == 200, response.text
    return response.json()


def publish(client, case, revision):
    preview = client.get(f'{API}/offerings/{case.offering}', headers=case.headers).json()['draft']['preview_hash']
    return client.post(f'{API}/offerings/{case.offering}/publish', headers=case.headers,
                       json={'expected_revision': revision, 'expected_preview_hash': preview, 'reason': 'Confirmed end of semester'})


def single_config(columns):
    return [{'key': 'regular', 'title': 'Assessment', 'weight': '100', 'columns': columns}]


def column(key='main', weight='100', **extra):
    return {'key': key, 'title': key, 'weight': weight, 'source_kind': 'manual', **extra}


def scored(case, values):
    return [{'student_id': student.id, 'column_key': key, 'state': 'scored', 'score': value}
            for student in (case.student, case.second) for key, value in values.items()]


def test_default_weights_pending_gate_and_exact_weight_validation(client, case):
    response = client.get(f'{API}/offerings/{case.offering}', headers=case.headers)
    # 批次 4 起新课程模板为平时 20 / 实验 30 / 期中 20 / 期末 30（各栏目默认等权）。
    assert [(c['key'], c['weight'], c['weight_mode']) for c in response.json()['config']] == \
           [('regular', '20', 'equal'), ('labs', '30', 'equal'), ('midterm', '20', 'equal'), ('final', '30', 'equal')]
    assert response.json()['revision'] == 0
    saved = save(client, case)
    assert len(saved['draft']['blockers']) == 8
    failed = publish(client, case, saved['revision'])
    assert failed.status_code == 409 and failed.json()['detail']['code'] == 'GRADEBOOK_INCOMPLETE'
    invalid = deepcopy(default_config()); invalid[0]['weight'] = '30.0001'
    assert client.put(f'{API}/offerings/{case.offering}/config', headers=case.headers,
                      json={'expected_revision': 1, 'categories': invalid}).status_code == 422
    assert client.get(f'{API}/offerings/{case.offering}', headers=case.headers).json()['revision'] == 1


def test_decimal_exemption_reweights_with_explicit_missing_and_complete_category_gate(client, case):
    save(client, case, single_config([column('a', '30'), column('b', '70')]))
    entries = scored(case, {'a': '33.33', 'b': '66.67'})
    # First student exempt a: only b retains 100% of the category.
    entries[0] = {'student_id': case.student.id, 'column_key': 'a', 'state': 'exempt', 'reason': 'Approved exemption'}
    entries[2] = {'student_id': case.second.id, 'column_key': 'a', 'state': 'missing', 'reason': 'Confirmed no submission'}
    edited = edit(client, case, entries, 1)
    rows = {r['student_id']: r for r in edited['draft']['rows']}
    assert rows[case.student.id]['final_score'] == '66.67'
    assert rows[case.student.id]['categories'][0]['effective_weights'] == {'b': '100.0000'}
    assert rows[case.second.id]['final_score'] == '46.67'
    assert publish(client, case, 2).status_code == 201
    current = client.get(f'{API}/offerings/{case.offering}', headers=case.headers).json()
    edited = edit(client, case, [{'student_id': case.student.id, 'column_key': 'b', 'state': 'exempt', 'reason': 'Second exemption'}], current['revision'])
    assert publish(client, case, edited['revision']).status_code == 409


def test_concurrent_revision_rejects_stale_edits_without_changing_scores(client, case):
    save(client, case, single_config([column()]))
    edit(client, case, scored(case, {'main': '80'}), 1)
    response = client.patch(f'{API}/offerings/{case.offering}/entries', headers=case.headers,
        json={'expected_revision': 1, 'entries': [{'student_id': case.student.id, 'column_key': 'main', 'state': 'scored', 'score': '99'}]})
    assert response.status_code == 409
    current = client.get(f'{API}/offerings/{case.offering}', headers=case.headers).json()
    assert current['revision'] == 2 and current['draft']['rows'][0]['final_score'] == '80.00'


def test_snapshots_archive_and_adjustment_are_immutable_and_student_filtered(client, db_session_factory, case):
    save(client, case, single_config([column()]))
    edit(client, case, scored(case, {'main': '80'}), 1)
    first = publish(client, case, 2).json()
    archived = client.post(f'{API}/offerings/{case.offering}/archive', headers=case.headers,
        json={'expected_revision': 3, 'reason': 'Archive original transcript'})
    assert archived.status_code == 201, archived.text
    archive_id = archived.json()['id']
    adjusted = client.post(f'{API}/offerings/{case.offering}/adjustment', headers=case.headers,
        json={'expected_revision': 4, 'student_id': case.student.id, 'score': '90', 'reason': 'Approved correction'})
    assert adjusted.status_code == 200
    second = publish(client, case, 5).json()
    assert second['payload']['rows'][0]['final_score'] == '90.00'
    original = client.get(f'{API}/snapshots/{archive_id}', headers=case.headers).json()
    assert original['payload']['rows'][0]['final_score'] == '80.00'
    assert original['payload_hash'] == first['payload_hash']
    own = client.get(f"{API}/snapshots/{second['id']}", headers=case.student_headers)
    assert len(own.json()['payload']['rows']) == 1 and 'Other Student Private' not in own.text
    for teacher in (case.other, case.owner):
        assert client.get(f'{API}/snapshots/{archive_id}', headers=auth_header(login(client, teacher.username)[0])).status_code == 403
    # Transfer/removal does not erase the student's own historical transcript.
    with db_session_factory() as db:
        row = db.scalar(select(OfferingEnrollment).where(OfferingEnrollment.student_id == case.student.id, OfferingEnrollment.offering_id == case.offering))
        row.status = 'transferred'; db.commit()
    assert client.get(f'{API}/snapshots/{archive_id}', headers=case.student_headers).status_code == 200
    exported = client.get(f'{API}/snapshots/{archive_id}/export', headers=case.student_headers)
    book = load_workbook(BytesIO(exported.content))
    assert book.active.max_row == 2 and book.active['B2'].data_type == 's'
    csv = client.get(f'{API}/snapshots/{archive_id}/export?format=csv', headers=case.student_headers)
    assert "'=Student" in csv.text and 'Other Student Private' not in csv.text
    # A later publication without the transferred student must keep the last own result discoverable.
    edit(client, case, [{'student_id': case.second.id, 'column_key': 'main', 'state': 'scored', 'score': '70'}], 6)
    assert publish(client, case, 7).status_code == 201
    own_history = client.get(f'{API}/offerings/{case.offering}', headers=case.student_headers)
    assert own_history.status_code == 200
    assert own_history.json()['snapshot']['id'] == second['id']


def exam_source(factory, case, offering_id=None, public=True):
    with factory() as db:
        exam = Exam(course_id=case.course, offering_id=offering_id or case.offering, title='Exam', status='published', show_score_after_grading=public)
        db.add(exam); db.flush()
        for points in (60, 140): db.add(ExamQuestion(exam_id=exam.id, question_type='fill_blank', prompt='Q', points=points, correct_answer={}))
        sub = ExamSubmission(exam_id=exam.id, student_id=case.student.id, status='graded', score=140)
        db.add(sub); db.commit()
        return exam.id, sub.id


def test_exam_percentage_publication_and_source_changes_leave_published_grade_fixed(client, db_session_factory, case):
    exam_id, submission_id = exam_source(db_session_factory, case)
    save(client, case, single_config([column(source_kind='exam', source_id=exam_id)]))
    assert publish(client, case, 1).status_code == 409  # Second student absence needs confirmation.
    edit(client, case, [{'student_id': case.second.id, 'column_key': 'main', 'state': 'missing', 'reason': 'Did not attend'}], 1)
    first = publish(client, case, 2)
    assert first.status_code == 201, first.text
    data = first.json(); assert data['payload']['rows'][0]['final_score'] == '70.00'
    with db_session_factory() as db:
        db.get(ExamSubmission, submission_id).score = 200; db.commit()
    current = client.get(f'{API}/offerings/{case.offering}', headers=case.headers).json()
    assert current['draft']['rows'][0]['final_score'] == '100.00'
    assert current['snapshot']['payload']['rows'][0]['final_score'] == '70.00'
    with db_session_factory() as db:
        assert db.get(GradebookSnapshot, data['id']).payload['rows'][0]['cells']['main']['source']['max_score'] == '200.0'


def test_unreleased_exam_and_wrong_offering_sources_do_not_publish(client, db_session_factory, case):
    private_id, _ = exam_source(db_session_factory, case, public=False)
    saved = save(client, case, single_config([column(source_kind='exam', source_id=private_id)]))
    assert saved['draft']['rows'][0]['cells']['main']['message'] == '考试成绩尚未公开'
    assert publish(client, case, 1).status_code == 409
    alien_id, _ = exam_source(db_session_factory, case, offering_id=case.other_offering)
    denied = client.put(f'{API}/offerings/{case.offering}/config', headers=case.headers,
        json={'expected_revision': 1, 'categories': single_config([column(source_kind='exam', source_id=alien_id)])})
    assert denied.status_code == 422
    assert client.get(f'{API}/offerings/{case.offering}', headers=case.headers).json()['revision'] == 1


def test_duplicate_sources_cross_student_edits_and_override_without_reason_rejected(client, db_session_factory, case):
    ident, _ = exam_source(db_session_factory, case)
    duplicate = single_config([column('one', '50', source_kind='exam', source_id=ident), column('two', '50', source_kind='exam', source_id=ident)])
    assert client.put(f'{API}/offerings/{case.offering}/config', headers=case.headers,
                      json={'expected_revision': 0, 'categories': duplicate}).status_code == 422
    save(client, case, single_config([column(source_kind='exam', source_id=ident)]))
    missing_reason = client.patch(f'{API}/offerings/{case.offering}/entries', headers=case.headers,
        json={'expected_revision': 1, 'entries': [{'student_id': case.student.id, 'column_key': 'main', 'state': 'scored', 'score': '100'}]})
    assert missing_reason.status_code == 422
    outsider = create_user(db_session_factory, 'outside-gradebook', 'student')
    outside = client.patch(f'{API}/offerings/{case.offering}/entries', headers=case.headers,
        json={'expected_revision': 1, 'entries': [{'student_id': outsider.id, 'column_key': 'main', 'state': 'scored', 'score': '90', 'reason': 'Cannot cross class'}]})
    assert outside.status_code == 422
    assert client.patch(f'{API}/offerings/{case.offering}/entries', headers=case.student_headers,
        json={'expected_revision': 1, 'entries': scored(case, {'main': '100'})}).status_code == 403
    with db_session_factory() as db:
        db.get(User, case.second.id).role = 'teacher'; db.commit()
    changed_role = client.patch(f'{API}/offerings/{case.offering}/entries', headers=case.headers,
        json={'expected_revision': 1, 'entries': [{'student_id': case.second.id, 'column_key': 'main', 'state': 'scored', 'score': '90', 'reason': 'Role was changed'}]})
    assert changed_role.status_code == 422


def test_partial_assignment_missing_needs_confirmation_and_keeps_earned_points(client, db_session_factory, case):
    with db_session_factory() as db:
        task = Assignment(course_id=case.course, offering_id=case.offering, title='Two questions', status='published')
        db.add(task); db.flush()
        questions = [JudgeQuestion(assignment_id=task.id, title=f'Q{i}', function_name='f', hidden_tests='') for i in range(2)]
        db.add_all(questions); db.flush()
        db.add(Submission(question_id=questions[0].id, student_id=case.student.id, code='', status='accepted', grading_status='completed', score=80))
        db.commit(); ident = task.id
    save(client, case, single_config([column(source_kind='assignment', source_id=ident)]))
    assert publish(client, case, 1).status_code == 409
    edited = edit(client, case, [{'student_id': s.id, 'column_key': 'main', 'state': 'missing', 'reason': 'Verified missing questions'} for s in (case.student, case.second)], 1)
    assert edited['draft']['rows'][0]['final_score'] == '40.00'
    assert edited['draft']['rows'][1]['final_score'] == '0.00'
    assert publish(client, case, 2).status_code == 201


def test_unpublished_draft_cannot_be_archived_in_place(client, case):
    save(client, case, single_config([column()]))
    edit(client, case, scored(case, {'main': '80'}), 1)
    assert publish(client, case, 2).status_code == 201
    edit(client, case, [{'student_id': case.student.id, 'column_key': 'main', 'state': 'scored', 'score': '75'}], 3)
    rejected = client.post(f'{API}/offerings/{case.offering}/archive', headers=case.headers,
        json={'expected_revision': 4, 'reason': 'Attempt draft archive'})
    assert rejected.status_code == 409 and rejected.json()['detail']['code'] == 'GRADEBOOK_UNPUBLISHED_CHANGES'
    latest = client.get(f'{API}/offerings/{case.offering}', headers=case.headers).json()
    assert latest['revision'] == 4 and latest['snapshot']['payload']['rows'][0]['final_score'] == '80.00'


def test_lab_platform_selected_published_revision_wins_over_column_highest(client, db_session_factory, test_settings, case):
    from app.models import ExperimentRecord, ExperimentSubmission, LabGradeRevision, LabRecord, LabRun, LabSubmission
    from app.models.teaching import OfferingExperiment
    from app.services.time_utils import utc_now
    from lab_test_support import headers, make_case
    lab = make_case(db_session_factory, test_settings)
    with db_session_factory() as db:
        record = ExperimentRecord(module_id=lab.module_id, student_id=lab.student.id,
            offering_id=lab.offering_id, template_version_id=lab.version_id, environment_version_id=lab.environment_id)
        db.add(record); db.flush()
        subs = [ExperimentSubmission(record_id=record.id, attempt_number=i + 1, score=value, cells_snapshot={})
                for i, value in enumerate((90, 20))]
        db.add_all(subs); db.flush()
        for submission in subs:
            grade = LabGradeRevision(submission_id=submission.id, assessment_id=lab.assessment_id,
                                     source='automatic', total=submission.score, published_at=utc_now())
            db.add(grade); db.flush()
            db.add(LabSubmission(submission_id=submission.id, assessment_id=lab.assessment_id,
                                 submission_hash=str(submission.id).zfill(64), published_grade_id=grade.id))
        db.add(LabRecord(record_id=record.id, assessment_id=lab.assessment_id, selected_submission_id=subs[1].id))
        release = db.scalar(select(OfferingExperiment).where(OfferingExperiment.offering_id == lab.offering_id))
        release_id, selected_id, record_id = release.id, subs[1].id, record.id
        db.commit()
    special = SimpleNamespace(offering=lab.offering_id, headers=headers(lab.teacher, test_settings))
    saved = save(client, special, single_config([column(source_kind='experiment', source_id=release_id, selection='highest')]))
    cell = saved['draft']['rows'][0]['cells']['main']
    assert cell['score'] == '20.00' and cell['source']['submission_id'] == selected_id
    assert cell['source']['selection'] == 'platform_selected'
    first = publish(client, special, 1)
    assert first.status_code == 201
    first_id, first_hash = first.json()['id'], first.json()['payload_hash']
    with db_session_factory() as db:
        new = ExperimentSubmission(record_id=record_id, attempt_number=3, cells_snapshot={}, outputs_snapshot={})
        db.add(new); db.flush()
        run = LabRun(owner_id=lab.student.id, kind='evaluate', record_id=record_id,
            assessment_id=lab.assessment_id, submission_id=new.id, status='queued')
        db.add(run); db.flush()
        db.add(LabSubmission(submission_id=new.id, assessment_id=lab.assessment_id,
            submission_hash='f' * 64, current_run_id=run.id))
        db.commit(); run_id = run.id
    # The old selected score stays historical; it cannot hide new grading or
    # infrastructure failures and become an apparently complete final grade.
    for status in ('queued', 'running', 'recovering', 'failed', 'needs_review'):
        with db_session_factory() as db:
            db.get(LabRun, run_id).status = status
            db.commit()
        draft = client.get(f'{API}/offerings/{special.offering}', headers=special.headers).json()['draft']
        assert draft['rows'][0]['cells']['main']['state'] == 'pending'
        assert draft['rows'][0]['final_score'] is None and draft['blockers']
        rejected = publish(client, special, 2)
        assert rejected.status_code == 409 and rejected.json()['detail']['code'] == 'GRADEBOOK_INCOMPLETE'
    # An explicit teacher correction with a reason can resolve an infrastructure
    # review; it produces a new version and leaves the earlier official result.
    edit(client, special, [{'student_id': lab.student.id, 'column_key': 'main', 'state': 'scored',
        'score': '20', 'reason': 'Infrastructure reviewed; retained approved score'}], 2)
    assert publish(client, special, 3).status_code == 201
    original = client.get(f'{API}/snapshots/{first_id}', headers=special.headers).json()
    assert original['payload_hash'] == first_hash and original['payload']['rows'][0]['final_score'] == '20.00'


def test_source_regrade_after_preview_requires_a_fresh_human_review(client, db_session_factory, case):
    ident, submission_id = exam_source(db_session_factory, case)
    save(client, case, single_config([column(source_kind='exam', source_id=ident)]))
    preview = edit(client, case, [{'student_id': case.second.id, 'column_key': 'main', 'state': 'missing', 'reason': 'Verified absence'}], 1)
    with db_session_factory() as db:
        db.get(ExamSubmission, submission_id).score = 200; db.commit()
    changed = client.post(f'{API}/offerings/{case.offering}/publish', headers=case.headers,
        json={'expected_revision': 2, 'expected_preview_hash': preview['draft']['preview_hash'], 'reason': 'Old review'})
    assert changed.status_code == 409 and changed.json()['detail']['code'] == 'GRADEBOOK_PREVIEW_CHANGED'
    with db_session_factory() as db:
        assert db.scalar(select(Gradebook).where(Gradebook.offering_id == case.offering)).revision == 2
        assert not db.scalars(select(GradebookSnapshot)).all()
    assert publish(client, case, 2).status_code == 201


def test_closed_term_prevents_grade_changes_but_preserves_official_archiving(client, db_session_factory, case):
    save(client, case, single_config([column()]))
    edit(client, case, scored(case, {'main': '80'}), 1)
    assert publish(client, case, 2).status_code == 201
    with db_session_factory() as db:
        offering = db.get(CourseOffering, case.offering)
        db.get(AcademicTerm, offering.academic_term_id).status = 'closed'
        db.commit()
    current = client.get(f'{API}/offerings/{case.offering}', headers=case.headers).json()
    assert current['read_only'] is True
    refused = client.patch(f'{API}/offerings/{case.offering}/entries', headers=case.headers,
        json={'expected_revision': 3, 'entries': scored(case, {'main': '90'})})
    assert refused.status_code == 409
    archived = client.post(f'{API}/offerings/{case.offering}/archive', headers=case.headers,
        json={'expected_revision': 3, 'reason': 'Archive after term closed'})
    assert archived.status_code == 201
    assert archived.json()['payload']['rows'][0]['final_score'] == '80.00'
    repeated = client.post(f'{API}/offerings/{case.offering}/archive', headers=case.headers,
        json={'expected_revision': 4, 'reason': 'Repeated archive'})
    assert repeated.status_code == 409 and repeated.json()['detail']['code'] == 'GRADEBOOK_ALREADY_ARCHIVED'


def test_teacher_entered_export_headers_remain_text_in_excel_and_csv(client, case):
    title = '=HYPERLINK("https://example.invalid", "grade")'
    configured = save(client, case, single_config([column(title=title)]))
    edited = edit(client, case, scored(case, {'main': '80'}), configured['revision'])
    result = publish(client, case, edited['revision'])
    assert result.status_code == 201, result.text
    snapshot_id = result.json()['id']
    spreadsheet = client.get(f'{API}/snapshots/{snapshot_id}/export', headers=case.headers)
    assert spreadsheet.status_code == 200
    workbook = load_workbook(BytesIO(spreadsheet.content), data_only=False)
    assert workbook.active['C1'].value == title
    assert workbook.active['C1'].data_type == 's'
    exported = client.get(f'{API}/snapshots/{snapshot_id}/export?format=csv', headers=case.headers)
    assert exported.status_code == 200
    headers = next(csv.reader(StringIO(exported.content.decode('utf-8-sig'))))
    assert headers[2] == "'" + title
