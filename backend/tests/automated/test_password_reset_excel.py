from io import BytesIO

from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import select

from app.models import AdminBatch, Course, Exam, ExamSubmission, User
from app.security import verify_password
from conftest import auth_header, create_user, login
from test_account_batches import admin_headers, confirm, exported_passwords, run_worker  # noqa: F401

API = '/api/v1/admin/password-resets'


def excel(rows, headers=('学号', '姓名')):
    wb = Workbook()
    wb.active.append(list(headers))
    for row in rows:
        wb.active.append(row)
    out = BytesIO(); wb.save(out)
    return out.getvalue()


@pytest.fixture()
def students(db_session_factory):
    users = [create_user(db_session_factory, 'excel-' + str(i), role)
             for i, role in enumerate(['student', 'student', 'student', 'teacher'], 1)]
    with db_session_factory() as db:
        for i, user in enumerate(users, 1):
            row = db.get(User, user.id)
            row.student_no = f'000{i}'
            row.real_name = f'姓名{i}'
            if i == 3:
                row.status = 'disabled'
        db.commit()
    return users


def upload(client, headers, rows, columns=('学号', '姓名')):
    return client.post(API + '/preview-file', headers=headers,
                       files={'file': ('reset.xlsx', excel(rows, columns),
                                       'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})


def test_excel_preview_preserves_credentials_and_reports_duplicates_disabled_and_exam(client, db_session_factory, admin_headers, students):
    with db_session_factory() as db:
        teacher = db.get(User, students[3].id)
        course = Course(title='Exam fixture', teacher_id=teacher.id)
        db.add(course); db.flush()
        exam = Exam(course_id=course.id, title='In progress')
        db.add(exam); db.flush()
        db.add(ExamSubmission(exam_id=exam.id, student_id=students[1].id, status='started'))
        before = {u.id: (u.password_hash, u.credential_version) for u in db.scalars(select(User))}
        db.commit()
    response = upload(client, admin_headers, [['0001', '姓名1'], ['0001', '姓名1'], ['0002'], ['0003']])
    assert response.status_code == 200, response.text
    batch = response.json()
    assert batch['summary'] == {'selected': 4, 'eligible': 1, 'skipped': 3, 'errors': 0}
    detail = client.get('/api/v1/admin/batches/' + batch['id'], headers=admin_headers).json()
    assert [r['row_number'] for r in detail['items']] == [2, 3, 4, 5]
    assert [r['status'] for r in detail['items']] == ['pending', 'skipped', 'skipped', 'skipped']
    assert '正在考试' in detail['items'][2]['result']['message']
    assert detail['items'][2]['data']['real_name'] == '姓名2'
    with db_session_factory() as db:
        assert before == {u.id: (u.password_hash, u.credential_version) for u in db.scalars(select(User))}
        assert db.get(AdminBatch, batch['id']).status == 'preview'


@pytest.mark.parametrize('row', [['missing', '不存在'], ['0001', '姓名错误'], ['0004', '姓名4'], [1, '姓名1'], ['=1+1', '姓名1'], ['', '姓名1']])
def test_invalid_rows_block_entire_confirmation(client, db_session_factory, admin_headers, students, row):
    response = upload(client, admin_headers, [['0002', '姓名2'], row])
    assert response.status_code == 200, response.text
    batch = response.json()
    assert batch['summary']['errors'] == 1
    assert sum(batch['summary'][key] for key in ['eligible', 'skipped', 'errors']) == batch['summary']['selected']
    result = confirm(client, admin_headers, batch, current_password='Passw0rd!', reason='Fixture reset')
    assert result.status_code == 409, result.text
    assert result.json()['detail']['code'] == 'PREVIEW_ERRORS'
    with db_session_factory() as db:
        assert db.get(AdminBatch, batch['id']).status == 'preview'
        assert verify_password('Passw0rd!', db.get(User, students[1].id).password_hash)


def test_valid_file_resets_only_matching_student_once_using_existing_confirmation_and_worker(client, db_session_factory, test_settings, admin_headers, students):
    response = upload(client, admin_headers, [['0001'], ['0001']], ('学号',))
    assert response.status_code == 200, response.text
    batch = response.json()
    assert confirm(client, admin_headers, batch, current_password='wrong', reason='Fixture reset').status_code == 422
    assert confirm(client, admin_headers, batch, current_password='Passw0rd!', reason='Fixture reset').status_code == 202
    run_worker(db_session_factory, test_settings)
    password = exported_passwords(client, admin_headers, batch['id'])['0001']
    with db_session_factory() as db:
        matched = db.get(User, students[0].id)
        assert matched.must_change_password
        assert verify_password(password, matched.password_hash)
        assert verify_password('Passw0rd!', db.get(User, students[1].id).password_hash)
    result = client.get('/api/v1/admin/batches/' + batch['id'], headers=admin_headers).json()
    assert [row['status'] for row in result['items']] == ['succeeded', 'skipped']


def test_template_and_non_admin_boundaries(client, db_session_factory, admin_headers, students):
    response = client.get(API + '/template', headers=admin_headers)
    wb = load_workbook(BytesIO(response.content), read_only=True)
    try:
        assert next(wb.active.iter_rows(values_only=True)) == ('学号', '姓名')
    finally:
        wb.close()
    headers = auth_header(login(client, students[0].username)[0])
    assert upload(client, headers, [['0001']]).status_code == 403
    assert client.get(API + '/template', headers=headers).status_code == 403
    assert upload(client, admin_headers, [['0001', '姓名1', 'secret']], ('学号', '姓名', '密码')).status_code == 422
